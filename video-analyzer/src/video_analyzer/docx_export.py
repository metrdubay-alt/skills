"""Small Markdown-to-DOCX exporter for generated analysis documents."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document


def markdown_to_docx(markdown: str, output_path: Path) -> None:
    doc = Document()
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "---":
            doc.add_paragraph("")
            i += 1
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            text = line[level:].strip()
            doc.add_heading(_strip_inline_markdown(text), level=level)
            i += 1
            continue
        if _is_table_start(lines, i):
            rows: list[list[str]] = []
            headers = _split_table_row(lines[i])
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = _strip_inline_markdown(text)
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = _strip_inline_markdown(text)
            continue
        if line.startswith("- "):
            _add_inline_runs(doc.add_paragraph(style="List Bullet"), line[2:].strip())
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            _add_inline_runs(doc.add_paragraph(style="List Number"), numbered.group(1).strip())
            i += 1
            continue
        _add_inline_runs(doc.add_paragraph(), line)
        i += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def _is_table_start(lines: list[str], i: int) -> bool:
    if i + 1 >= len(lines):
        return False
    return (
        lines[i].strip().startswith("|")
        and lines[i + 1].strip().startswith("|")
        and set(lines[i + 1].replace("|", "").strip()) <= {"-", ":", " "}
    )


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _strip_inline_markdown(text: str) -> str:
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", re.sub(r"`([^`]+)`", r"\1", text))
